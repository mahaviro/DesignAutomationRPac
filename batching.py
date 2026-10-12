import pandas as pd
import streamlit as st
import openpyxl
from openpyxl.utils import column_index_from_string
from openpyxl.styles import Font, Alignment, Border, Side
import numpy as np
from datetime import datetime
import os
import json
import shutil

st.set_page_config(page_title="Print Plate Batching & Export Application", layout="wide")

st.title("Plate Batching Automation Application")

CONFIG_FILE = "batching_config.json"

def load_config():
    if os.path.exists(CONFIG_FILE):
        try:
            with open(CONFIG_FILE, "r") as f:
                return json.load(f)
        except:
            pass
    return {
        "input_method": "Upload File",
        "dir_path": "D:/Automation File Maker/",
        "filename": "order1.xlsx",
        "sheet_input_name": "",
        "desc_cell_input": "B1",
        "header_row": 22,
        "col_line": "A",
        "col_style": "E",
        "col_color": "F",
        "col_size": "H",
        "col_barcode": "AX",
        "col_qty": "DH",
        "col_item_code": "DI",
        "ups_per_impression": 36,
        "target_plates_input": 1,
        "overrun_allowance": 10.0,
        "mo_number": "260041976",
        "ref_name": "CARTER HT",
        "output_dir_reports": "./Batching_Outputs/",
        "output_dir_data": "./Data_Exports/"
    }

def save_config(cfg):
    try:
        with open(CONFIG_FILE, "w") as f:
            json.dump(cfg, f, indent=4)
    except Exception as e:
        st.warning(f"Failed to save configuration: {e}")

config = load_config()

with st.sidebar:
    st.header("⚙️ Configuration Menu")
    
    if st.button("🔄 Reset the Document", type="secondary", use_container_width=True):
        for key in list(st.session_state.keys()):
            del st.session_state[key]
        st.success("Session has been reset, please choose your document!")
        st.rerun()
    
    st.markdown("---")
    
    with st.expander("📁 1. Excel Data Source & Settings", expanded=False):
        input_method = st.radio("Input Method", ["Upload File", "Read from Local Directory"], index=0 if config["input_method"]=="Upload File" else 1)
        file_path_or_upload = None
        if input_method == "Upload File":
            file_path_or_upload = st.file_uploader("Select Order Excel File (.xlsx)", type=["xlsx"])
        else:
            dir_path = st.text_input("Directory Path", value=config["dir_path"])
            filename = st.text_input("Excel Filename", value=config["filename"])
            file_path_or_upload = dir_path + filename

        sheet_input_name = st.text_input("Excel Sheet Name (Leave blank for first sheet)", value=config["sheet_input_name"])
        desc_cell_input = st.text_input("Description Cell Coordinate (e.g., B1, B2, A2)", value=config["desc_cell_input"])

    with st.expander("⚙️ 2. Header & Column Mapping", expanded=False):
        header_row = st.number_input("Header Row (1-based index)", min_value=1, value=int(config["header_row"]), step=1)
        col_line = st.text_input("Line Number Column", value=config["col_line"])
        col_style = st.text_input("Style Column", value=config["col_style"])
        col_color = st.text_input("Color Column", value=config["col_color"])
        col_size = st.text_input("Size Column", value=config["col_size"])
        col_barcode = st.text_input("Barcode Column", value=config["col_barcode"])
        col_qty = st.text_input("Qty Column", value=config["col_qty"])
        col_item_code = st.text_input("Item Code Column (e.g., DI)", value=config["col_item_code"])

    with st.expander("⚙️ 3. Batching & Report Parameters", expanded=False):
        ups_per_impression = st.number_input("UPS per Impression (Max Ups/Plate / R8)", min_value=1, value=int(config["ups_per_impression"]), step=1)
        target_plates_input = st.number_input("Target Number of Plates (for Qty < 2500)", min_value=1, value=int(config["target_plates_input"]), step=1)
        overrun_allowance_input = st.number_input("Overrun Allowance (%)", min_value=0.0, max_value=100.0, value=float(config["overrun_allowance"]), step=0.5)
        overrun_allowance = overrun_allowance_input / 100.0
        
        mo_number = st.text_input("MO Number (MO#)", value=config["mo_number"])
        ref_name = st.text_input("Reference Name (REF#)", value=config["ref_name"])
        
    with st.expander("📂 4. Export Directories", expanded=False):
        output_dir_reports = st.text_input("Batching Formula & Report Folder", value=config["output_dir_reports"])
        output_dir_data = st.text_input("Excel Data Output Folder", value=config["output_dir_data"])

    if st.button("💾 Save Parameters Configuration", use_container_width=True):
        new_cfg = {
            "input_method": input_method,
            "dir_path": dir_path if input_method == "Read from Local Directory" else config["dir_path"],
            "filename": filename if input_method == "Read from Local Directory" else config["filename"],
            "sheet_input_name": sheet_input_name,
            "desc_cell_input": desc_cell_input,
            "header_row": header_row,
            "col_line": col_line,
            "col_style": col_style,
            "col_color": col_color,
            "col_size": col_size,
            "col_barcode": col_barcode,
            "col_qty": col_qty,
            "col_item_code": col_item_code,
            "ups_per_impression": ups_per_impression,
            "target_plates_input": target_plates_input,
            "overrun_allowance": overrun_allowance_input,
            "mo_number": mo_number,
            "ref_name": ref_name,
            "output_dir_reports": output_dir_reports,
            "output_dir_data": output_dir_data
        }
        save_config(new_cfg)
        st.success("✅ Configuration saved successfully!")

    # Watermark Footer di Sidebar
    st.markdown("---")
    st.markdown(
        "<div style='text-align: center; color: #888888; font-size: 12px; font-weight: bold; padding-bottom: 10px;'>"
        "PT. r-pac Packaging Indonesia © 2026"
        "</div>", 
        unsafe_allow_html=True
    )

if file_path_or_upload:
    try:
        wb_data = openpyxl.load_workbook(file_path_or_upload, data_only=True)
        sheet_names = wb_data.sheetnames
        selected_sheet = sheet_input_name if sheet_input_name in sheet_names else sheet_names[0]
        ws_data = wb_data[selected_sheet]
        
        desc_val = ""
        cell_coord = desc_cell_input.strip().upper()
        if cell_coord:
            try:
                desc_val = ws_data[cell_coord].value
                if desc_val is None:
                    wb_formula = openpyxl.load_workbook(file_path_or_upload, data_only=False)
                    desc_val = wb_formula[selected_sheet][cell_coord].value
            except Exception:
                pass
        
        if not desc_val:
            for fallback in ['B1', 'A2', 'B2', 'A1']:
                try:
                    v = ws_data[fallback].value
                    if v:
                        desc_val = v
                        break
                except:
                    pass

        desc_val = str(desc_val).strip() if desc_val is not None else ""
        
        df_raw = pd.read_excel(file_path_or_upload, sheet_name=selected_sheet, header=header_row-1)
        
        def get_col_name(df, letter):
            try:
                idx = column_index_from_string(letter.upper()) - 1
                if idx < len(df.columns):
                    return df.columns[idx]
            except:
                pass
            return None

        name_line = get_col_name(df_raw, col_line)
        name_style = get_col_name(df_raw, col_style)
        name_color = get_col_name(df_raw, col_color)
        name_size = get_col_name(df_raw, col_size)
        name_barcode = get_col_name(df_raw, col_barcode)
        name_qty = get_col_name(df_raw, col_qty)
        name_item_code = get_col_name(df_raw, col_item_code)

        st.subheader("📋 Reference Data Preview & Source Control")
        st.info(f"📂 File: `{file_path_or_upload}` | Sheet: `{selected_sheet}` | Header Row: `{header_row}`")
        st.success(f"📝 **Extracted Description (from Cell `{desc_cell_input}`):** `{desc_val}`")

        if all([name_line, name_style, name_color, name_size, name_barcode, name_qty]):
            raw_q = df_raw[name_qty]
            clean_q = pd.to_numeric(raw_q, errors='coerce').fillna(0).values if isinstance(raw_q, pd.Series) else [0.0]*len(df_raw)
            
            raw_ic = df_raw[name_item_code] if name_item_code else pd.Series(["DEFAULT"] * len(df_raw))
            clean_ic = raw_ic.fillna("DEFAULT").astype(str).values if isinstance(raw_ic, pd.Series) else ["DEFAULT"]*len(df_raw)

            df_preview = pd.DataFrame({
                'No': df_raw[name_line],
                'Item Code': clean_ic,
                'Description': [desc_val] * len(df_raw),
                'Style': df_raw[name_style],
                'Color': df_raw[name_color],
                'Size': df_raw[name_size],
                'Barcode': df_raw[name_barcode],
                'Qty': clean_q,
                'Extra Qty': 0.0
            })
            
            df_preview = df_preview.dropna(subset=['Barcode', 'Style']).reset_index(drop=True)
            
            column_config = {
                "No": st.column_config.NumberColumn("No", disabled=True),
                "Item Code": st.column_config.TextColumn("Item Code", disabled=False),
                "Barcode": st.column_config.TextColumn("Barcode", disabled=True),
                "Style": st.column_config.TextColumn("Style", disabled=True),
                "Size": st.column_config.TextColumn("Size", disabled=True),
                "Color": st.column_config.TextColumn("Color", disabled=True),
                "Description": st.column_config.TextColumn("Description", disabled=False),
                "Qty": st.column_config.NumberColumn("Qty", disabled=True),
                "Extra Qty": st.column_config.NumberColumn("Extra Qty", disabled=False, min_value=0, step=1)
            }
            
            edited_preview = st.data_editor(
                df_preview, 
                num_rows="fixed", 
                key="data_acuan_editor", 
                column_config=column_config, 
                use_container_width=True
            )
            
            st.markdown("---")
            st.subheader("🚀 Optimized Plate Batching & Export")
            
            if 'calc_results' not in st.session_state:
                st.session_state.calc_results = None

            def compute_batching(df_source, is_recalc=False, num_target_plates=1, force_recluster=False, threshold=2500):
                df = df_source.copy()
                q_s = pd.to_numeric(df['Qty'], errors='coerce').fillna(0)
                eq_s = pd.to_numeric(df['Extra Qty'], errors='coerce').fillna(0)
                
                df['Qty'] = q_s.values
                df['Extra Qty'] = eq_s.values
                df['TOTAL (Pcs)'] = df['Qty'] + df['Extra Qty']
                df = df[df['TOTAL (Pcs)'] > 0].reset_index(drop=True)
                
                if len(df) == 0:
                    return df
                
                need_auto_plate = ('PLATE' not in df.columns) or (not is_recalc) or force_recluster
                
                if need_auto_plate:
                    df['Original_Index'] = df.index
                    
                    is_dedicated = df['TOTAL (Pcs)'] >= threshold
                    df_dedicated = df[is_dedicated].copy()
                    df_regular = df[~is_dedicated].copy()
                    
                    plate_assignments = {}
                    next_plate_id = 1
                    
                    for idx in df_dedicated['Original_Index']:
                        plate_assignments[idx] = next_plate_id
                        next_plate_id += 1
                        
                    if len(df_regular) > 0:
                        effective_target_plates = min(num_target_plates, len(df_regular))
                        total_qty_sum = df_regular['TOTAL (Pcs)'].sum()
                        if total_qty_sum > 0:
                            df_regular['Raw_UPS'] = (df_regular['TOTAL (Pcs)'] / total_qty_sum) * ups_per_impression
                            df_regular['Temp_UPS'] = df_regular['Raw_UPS'].apply(lambda x: max(1, int(round(x))))
                            df_regular['Temp_Sheets'] = df_regular.apply(lambda r: int(np.ceil(r['TOTAL (Pcs)'] / r['Temp_UPS'])) if r['Temp_UPS'] > 0 else 1, axis=1)
                            
                            df_sorted = df_regular.sort_values(by=['Temp_Sheets', 'TOTAL (Pcs)'], ascending=[False, False]).reset_index(drop=True)
                            
                            plate_capacities = {}
                            total_items = len(df_sorted)
                        
                            for idx_s, row in df_sorted.iterrows():
                                u = row['Temp_UPS']
                                
                                target_p = ((idx_s * effective_target_plates) // total_items) + 1
                                target_p = min(effective_target_plates, max(1, target_p))
                                
                                if target_p not in plate_capacities:
                                    plate_capacities[target_p] = 0
                                    
                                if plate_capacities[target_p] + u > ups_per_impression:
                                    assigned_plate = None
                                    for p_id in sorted(plate_capacities.keys()):
                                        if plate_capacities[p_id] + u <= ups_per_impression:
                                            assigned_plate = p_id
                                            plate_capacities[p_id] += u
                                            break
                                    if assigned_plate is None:
                                        new_p_id = len(plate_capacities) + 1
                                        plate_capacities[new_p_id] = u
                                        assigned_plate = new_p_id
                                    reg_plate_num = assigned_plate + next_plate_id - 1
                                else:
                                    plate_capacities[target_p] += u
                                    reg_plate_num = target_p + next_plate_id - 1
                                    
                                plate_assignments[row['Original_Index']] = reg_plate_num
                                
                    df['PLATE'] = df['Original_Index'].map(plate_assignments)
                    df = df.sort_values(by='Original_Index').reset_index(drop=True)
                
                plate_s = pd.to_numeric(df['PLATE'], errors='coerce')
                df['PLATE'] = plate_s.fillna(1).astype(int).values if isinstance(plate_s, pd.Series) else 1
                
                adj_s = pd.to_numeric(df.get('ADJUST UPS', 0), errors='coerce')
                df['ADJUST UPS'] = adj_s.fillna(0).astype(int).values if isinstance(adj_s, pd.Series) else 0
                
                for p_id, group in df.groupby('PLATE'):
                    total_plate_qty = group['TOTAL (Pcs)'].sum()
                    indices = group.index
                    if total_plate_qty > 0:
                        plate_ups_list = []
                        for idx in indices:
                            q_item = df.loc[idx, 'TOTAL (Pcs)']
                            adj_item = df.loc[idx, 'ADJUST UPS']
                            raw_share = (q_item / total_plate_qty) * ups_per_impression
                            assigned_u = max(1, int(round(raw_share)) + adj_item)
                            plate_ups_list.append((idx, assigned_u))
                        
                        current_plate_ups_sum = sum([u for _, u in plate_ups_list])
                        diff = ups_per_impression - current_plate_ups_sum
                        if plate_ups_list:
                            first_idx, first_u = plate_ups_list[0]
                            new_first_u = max(1, first_u + diff)
                            plate_ups_list[0] = (first_idx, new_first_u)
                        
                        for idx, u_val in plate_ups_list:
                            df.loc[idx, 'UPS'] = u_val

                df['UPS'] = df['UPS'].apply(lambda x: max(1, int(x)))
                df['Total UPS'] = df.groupby('PLATE')['UPS'].transform('sum')
                df['Sheets Actual'] = df.apply(lambda r: int(np.ceil(r['TOTAL (Pcs)'] / r['UPS'])) if r['UPS'] > 0 else 0, axis=1)
                
                max_sheets_per_plate = df.groupby('PLATE')['Sheets Actual'].transform('max')
                df['Net Sheets'] = max_sheets_per_plate
                df['Sheets'] = df['Net Sheets']
                
                df['WT'] = df['Description']
                df['Helper'] = ((df.duplicated(subset=['WT'], keep='first') == False).astype(int)) if 'WT' in df.columns else 1
                df['Plate (U)'] = df['PLATE']
                df['Print Qty'] = df['Net Sheets'] * df['UPS']
                df['Over run Pcs'] = (df['Print Qty'] - df['TOTAL (Pcs)']).apply(lambda x: max(0, x))
                
                for idx, row in df.iterrows():
                    if row['Print Qty'] < row['TOTAL (Pcs)'] and row['UPS'] > 0:
                        needed_sheets = np.ceil(row['TOTAL (Pcs)'] / row['UPS'])
                        df.loc[idx, 'Net Sheets'] = needed_sheets
                        df.loc[idx, 'Sheets'] = needed_sheets
                        df.loc[idx, 'Print Qty'] = needed_sheets * row['UPS']
                        df.loc[idx, 'Over run Pcs'] = max(0, df.loc[idx, 'Print Qty'] - row['TOTAL (Pcs)'])
                
                df['Size wise over run %'] = df.apply(lambda r: f"{(r['Over run Pcs'] / r['TOTAL (Pcs)'] * 100):.2f}%" if r['TOTAL (Pcs)'] > 0 else "0.00%", axis=1)
                df['Batching Scrap'] = 0.0
                df['Over run Sheet'] = df.apply(lambda r: r['Over run Pcs'] / ups_per_impression if ups_per_impression > 0 else 0, axis=1)
                
                return df

            def find_best_optimal_plates(df_preview_data):
                q_s = pd.to_numeric(df_preview_data['Qty'], errors='coerce').fillna(0)
                eq_s = pd.to_numeric(df_preview_data['Extra Qty'], errors='coerce').fillna(0)
                total_pcs = q_s + eq_s
                is_ded = total_pcs >= 2500
                reg_count = len(total_pcs[~is_ded])
                
                if reg_count <= 1:
                    return max(1, reg_count)
                
                best_p = 1
                min_overrun_pcs = float('inf')
                
                max_test = min(reg_count, 15)
                for p in range(1, max_test + 1):
                    res_df = compute_batching(df_preview_data, is_recalc=False, num_target_plates=p, threshold=2500)
                    total_overrun = res_df['Over run Pcs'].sum()
                    if total_overrun < min_overrun_pcs:
                        min_overrun_pcs = total_overrun
                        best_p = p
                return best_p

            col_btn1, col_btn2 = st.columns(2)
            run_calc = col_btn1.button("Run Batching Calculation", type="primary", use_container_width=True)
            recalc = col_btn2.button("🔄 Recalculate Batching", type="secondary", use_container_width=True)
            
            # Run pertama kali otomatis mencari jumlah plat terbaik
            if st.session_state.calc_results is None and not run_calc and not recalc:
                optimal_plates = find_best_optimal_plates(edited_preview)
                st.session_state.calc_results = compute_batching(edited_preview, is_recalc=False, num_target_plates=optimal_plates, threshold=2500)

            if run_calc:
                optimal_plates = find_best_optimal_plates(edited_preview)
                st.session_state.calc_results = compute_batching(edited_preview, is_recalc=False, num_target_plates=optimal_plates, threshold=2500)
                st.success(f"✅ Calculation completed successfully using automatically optimized optimal plates ({optimal_plates} plates)!")

            if st.session_state.calc_results is not None:
                output_columns = [
                    'No', 'Item Code', 'Barcode', 'Style', 'Size', 'Qty', 'Extra Qty', 'Color', 'Description', 
                    'TOTAL (Pcs)', 'PLATE', 'UPS', 'Total UPS', 'Sheets', 'WT', 'Sheets Actual', 
                    'ADJUST UPS', 'Plate (U)', 'Net Sheets', 'Print Qty', 'Over run Pcs', 
                    'Size wise over run %', 'Batching Scrap', 'Helper', 'Over run Sheet'
                ]
                
                st.markdown("### 📊 Batching Results Table (Fully Recalculable & Editable)")
                
                final_batch_editor = st.data_editor(
                    st.session_state.calc_results[output_columns],
                    key="final_batch_editor",
                    use_container_width=True
                )

                if recalc:
                    updated_df = final_batch_editor.copy()
                    st.session_state.calc_results = compute_batching(updated_df, is_recalc=True, num_target_plates=target_plates_input, force_recluster=True, threshold=2500)
                    st.success(f"🔄 Batching successfully recalculated using manual target {target_plates_input} plates from configuration!")
                    st.rerun()

                total_plates = final_batch_editor['PLATE'].nunique()
                overall_plate = final_batch_editor['Helper'].sum()
                total_print_pcs = final_batch_editor['Print Qty'].sum()
                total_order_pcs = final_batch_editor['Qty'].sum()
                total_extra_pcs = final_batch_editor['Extra Qty'].sum()
                total_overrun_pcs = final_batch_editor['Over run Pcs'].sum()
                total_size_wise_overrun_pct = (total_overrun_pcs / total_order_pcs * 100) if total_order_pcs > 0 else 0.0
                
                c1, c2, c3, c4, c5, c6, c7 = st.columns(7)
                c1.metric("Total Qty Order", f"{total_order_pcs:,.0f}")
                c2.metric("Total Extra Qty", f"{total_extra_pcs:,.0f}")
                c3.metric("Total Plates Used", total_plates)
                c4.metric("Overall Plate", overall_plate)
                c5.metric("Total Over run Pcs", f"{total_overrun_pcs:,.0f}")
                c6.metric("Total Print Pcs", f"{total_print_pcs:,}")
                c7.metric("Total Size-wise Over %", f"{total_size_wise_overrun_pct:.2f}%")
                
                st.markdown("---")
                
                if st.button("💾 Export All Files (Formula, Report & Excel Data)", type="primary", use_container_width=True):
                    try:
                        base_mo_folder_path = os.path.join(output_dir_reports, str(mo_number))
                        mo_folder_path = base_mo_folder_path
                        counter = 1
                        while os.path.exists(mo_folder_path) and any(os.scandir(mo_folder_path)):
                            mo_folder_path = f"{base_mo_folder_path}_{counter}"
                            counter += 1
                        
                        os.makedirs(mo_folder_path, exist_ok=True)
                        os.makedirs(output_dir_data, exist_ok=True)
                        
                        item_code_val = str(final_batch_editor['Item Code'].iloc[0]).strip() if 'Item Code' in final_batch_editor.columns and len(final_batch_editor) > 0 else "ITEM"
                        
                        thin_side = Side(border_style="thin", color="000000")
                        border_all = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
                        align_center = Alignment(horizontal="center", vertical="center")
                        align_left = Alignment(horizontal="left", vertical="center")
                        align_right = Alignment(horizontal="right", vertical="center")
                        
                        num_rows = len(final_batch_editor)
                        end_r = 11 + num_rows
                        
                        # 1. Export Batching Formula Excel (Pcs-only summary)
                        formula_filename = os.path.join(mo_folder_path, f"{item_code_val}_{mo_number}_Batching_Formula.xlsx")
                        wb_f = openpyxl.Workbook()
                        ws_f = wb_f.active
                        ws_f.title = "BATCHING"
                        
                        ws_f['B2'] = "BATCHING MASTER"
                        ws_f['B2'].font = Font(name="Arial", size=14, bold=True)
                        ws_f['B4'] = "DATE :"
                        ws_f['C4'] = datetime.now()
                        ws_f['C4'].number_format = 'DD-MM-YY'
                        ws_f['B5'] = "REF :"
                        ws_f['C5'] = ref_name
                        ws_f['B6'] = "MO#   :"
                        ws_f['C6'] = mo_number
                        
                        ws_f['E7'] = "TOTAL QTY ORDER"
                        ws_f['G7'] = "TOTAL QTY EXTRA"
                        ws_f['I7'] = "TOTAL QTY (Pcs)"
                        ws_f['K7'] = "TOTAL QTY PRINT"
                        ws_f['M7'] = "TOTAL OVER RUN"
                        ws_f['O7'] = "TOTAL OVERRUN %"
                        
                        ws_f[f'E8'] = f"=SUM(F12:F{end_r})"
                        ws_f[f'G8'] = f"=SUM(G12:G{end_r})"
                        ws_f[f'I8'] = f"=SUM(J12:J{end_r})"
                        ws_f[f'K8'] = f"=SUM(V12:V{end_r})"
                        ws_f[f'M8'] = f"=SUM(W12:W{end_r})"
                        ws_f[f'O8'] = f"=M8/E8"
                        ws_f[f'O8'].number_format = '0.00%'
                        
                        for col_let in ['E', 'G', 'I', 'K', 'M', 'O']:
                            ws_f[f'{col_let}7'].font = Font(name="Arial", size=9, bold=True)
                            ws_f[f'{col_let}7'].alignment = align_center
                            ws_f[f'{col_let}8'].font = Font(name="Arial", size=10, bold=True)
                            ws_f[f'{col_let}8'].alignment = align_center

                        headers_formula = ['No', 'Barcode', 'Style', 'Size', 'Qty', 'Extra Qty', 'Color', 'Description', 'TOTAL (Pcs)', 'AVG', 'PLATE', 'Ups Per Size', 'UPS', 'Total UPS', 'Sheets', 'WT', 'Sheets Actual', 'ADJUST UPS', 'Plate', 'Net Sheets', 'Print Qty', 'Over run Pcs', 'Size wise over run %', 'Batching Scrap', 'Helper', 'Over run Sheet']
                        ws_f.cell(row=10, column=2, value="Req Ups Qty & Sheets Qty")
                        ws_f.cell(row=10, column=2).font = Font(name="Arial", size=11, bold=True)
                        
                        ws_f.row_dimensions[11].height = 28
                        ws_f.cell(row=11, column=2, value="No")
                        for idx_h, h_name in enumerate(headers_formula):
                            c_cell = ws_f.cell(row=11, column=idx_h + 2, value=h_name)
                            c_cell.font = Font(name="Arial", size=10, bold=True)
                            c_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                            c_cell.border = border_all
                        
                        for idx, row in final_batch_editor.iterrows():
                            r_num = 12 + idx
                            f_qty = row.get('Qty')
                            f_ext = row.get('Extra Qty')
                            f_plate = row.get('PLATE')
                            f_ups = row.get('UPS')
                            f_netsheets = row.get('Net Sheets')
                            
                            row_data = [
                                None,
                                row.get('No'), str(row.get('Barcode', '')), row.get('Style'), row.get('Size'),
                                f_qty, f_ext, row.get('Color'), row.get('Description'),
                                f"=IF(F{r_num}=\"\",\"\",F{r_num}+G{r_num})",
                                f"=IFERROR(ROUND(J{r_num}/SUM($J$12:$J${end_r})*{ups_per_impression},0),\"\")",
                                f_plate,
                                f"=IFERROR(J{r_num}/P{r_num},0)",
                                f_ups,
                                f"=IF(C{r_num}=\"\",\"\",SUMIFS($N$12:$N${end_r},$L$12:$L${end_r},L{r_num}))",
                                f"=SUMIF($L$12:$L${end_r},L{r_num},$J$12:$J${end_r})/$R$8",
                                f"=IF(I{r_num}<>\"\",$C$6,\"\")",
                                f"=ROUNDUP(IF(J{r_num}=\"\",0,J{r_num}/N{r_num}),0)",
                                f"=L{r_num}",
                                f"=L{r_num}",
                                f_netsheets,
                                f"=U{r_num}*N{r_num}",
                                f"=V{r_num}-J{r_num}",
                                f"=IF(J{r_num}=\"\",\"\",W{r_num}/J{r_num})",
                                0,
                                f"=IF(COUNTIF($Q$12:$Q{r_num},$Q{r_num})>1,\"\",1)",
                                f"=W{r_num}/{ups_per_impression}"
                            ]
                            ws_f.append(row_data)
                            curr_r = ws_f.max_row
                            ws_f.row_dimensions[curr_r].height = 20
                            for c_idx in range(2, len(row_data) + 1):
                                c_cell = ws_f.cell(row=curr_r, column=c_idx)
                                c_cell.font = Font(name="Arial", size=10)
                                c_cell.border = border_all
                                c_cell.alignment = align_center
                                if c_idx == 24:
                                    c_cell.number_format = '0.00%'
                        
                        for col in ws_f.columns:
                            max_len = max(len(str(cell.value or '')) for cell in col)
                            col_letter = openpyxl.utils.get_column_letter(col[0].column)
                            ws_f.column_dimensions[col_letter].width = max(max_len + 3, 10)

                        wb_f.save(formula_filename)
                        
                        # 2. Export Batching Report Excel
                        report_filename = os.path.join(mo_folder_path, f"{item_code_val}_{mo_number}_Batching_Report.xlsx")
                        wb_r = openpyxl.Workbook()
                        ws_r = wb_r.active
                        ws_r.title = str(mo_number)
                        
                        ws_r['A1'] = "MO#"
                        ws_r['A1'].font = Font(name="Arial", size=16, bold=True)
                        ws_r['B1'] = mo_number
                        ws_r['B1'].font = Font(name="Arial", size=18, bold=True)
                        
                        ws_r['A2'] = "REF#"
                        ws_r['A2'].font = Font(name="Arial", size=16, bold=True)
                        ws_r['B2'] = ref_name
                        ws_r['B2'].font = Font(name="Arial", size=18, bold=True)
                        
                        ws_r['A3'] = "DATE :"
                        ws_r['A3'].font = Font(name="Arial", size=14, bold=True)
                        ws_r['B3'] = datetime.now()
                        ws_r['B3'].number_format = 'DD-MM-YY'
                        ws_r['B3'].font = Font(name="Arial", size=14)
                        
                        headers_report = ['LINE#', 'Barcode', 'Style', 'Size', 'Qty', 'Extra Qty', 'Color', 'Description', 'TOTAL QTY (Pcs)', 'PLATE', 'UPS', 'NEED SHEET', 'PRINT QTY', 'OVER RUN PCS', 'REMARK']
                        ws_r.append([]) # row 4
                        ws_r.append(headers_report) # row 5
                        ws_r.row_dimensions[5].height = 25
                        
                        for col_idx in range(1, len(headers_report) + 1):
                            cell = ws_r.cell(row=5, column=col_idx)
                            cell.font = Font(name="Arial", size=11, bold=True)
                            cell.alignment = align_center
                            cell.border = border_all
                        
                        start_row = 6
                        for idx, row in final_batch_editor.iterrows():
                            r_data = [
                                row.get('No'), str(row.get('Barcode', '')), row.get('Style'), row.get('Size'),
                                row.get('Qty'), row.get('Extra Qty'), row.get('Color'), row.get('Description'),
                                row.get('TOTAL (Pcs)'), row.get('PLATE'), row.get('UPS'), row.get('Net Sheets'),
                                row.get('Print Qty'), row.get('Over run Pcs'), None
                            ]
                            ws_r.append(r_data)
                            curr_r = ws_r.max_row
                            ws_r.row_dimensions[curr_r].height = 20
                            for c_idx in range(1, len(r_data) + 1):
                                c_cell = ws_r.cell(row=curr_r, column=c_idx)
                                c_cell.font = Font(name="Arial", size=11)
                                c_cell.border = border_all
                                if c_idx in [1, 4, 5, 6, 7, 10, 11, 12, 13, 14]:
                                    c_cell.alignment = align_center
                                elif c_idx in [2, 3, 8, 15]:
                                    c_cell.alignment = align_left
                                elif c_idx == 9:
                                    c_cell.alignment = align_right
                                    c_cell.font = Font(name="Arial", size=12, bold=True)
                        
                        total_row_idx = start_row + len(final_batch_editor) - 1
                        tot_data = [None, None, None, None, f"=SUM(E{start_row}:E{total_row_idx})", f"=SUM(F{start_row}:F{total_row_idx})", None, 'TOTAL QTY (PCS)', f"=SUM(I{start_row}:I{total_row_idx})", None, None, None, f"=SUM(M{start_row}:M{total_row_idx})", f"=SUM(N{start_row}:N{total_row_idx})", None]
                        ws_r.append(tot_data)
                        tot_r = ws_r.max_row
                        ws_r.row_dimensions[tot_r].height = 22
                        
                        for c_idx in range(1, len(tot_data) + 1):
                            c_cell = ws_r.cell(row=tot_r, column=c_idx)
                            c_cell.font = Font(name="Arial", size=12, bold=True)
                            if c_idx in [5, 6, 8, 9, 13, 14]:
                                c_cell.border = border_all
                                if c_idx in [5, 6, 9, 13, 14]:
                                    c_cell.alignment = align_center
                                elif c_idx == 8:
                                    c_cell.alignment = align_right

                        wb_r.save(report_filename)
                        
                        # 3. Export Excel Data (Preserving original file template, metadata rows 1-21, and adding Plate & UPS columns)
                        data_filename = os.path.join(output_dir_data, f"{item_code_val}_Data.xlsx")
                        wb_orig_export = openpyxl.load_workbook(file_path_or_upload)
                        ws_orig_export = wb_orig_export[selected_sheet]
                        
                        max_c = ws_orig_export.max_column
                        ws_orig_export.cell(row=header_row, column=max_c + 1, value="Plate")
                        ws_orig_export.cell(row=header_row, column=max_c + 2, value="UPS")
                        
                        line_col_idx = column_index_from_string(col_line.upper())
                        ups_map = dict(zip(final_batch_editor['No'], final_batch_editor['UPS']))
                        plate_map = dict(zip(final_batch_editor['No'], final_batch_editor['PLATE']))
                        
                        for r in range(header_row + 1, ws_orig_export.max_row + 1):
                            line_val = ws_orig_export.cell(row=r, column=line_col_idx).value
                            if line_val is not None:
                                try:
                                    line_int = int(line_val)
                                    if line_int in ups_map:
                                        ws_orig_export.cell(row=r, column=max_c + 1, value=plate_map[line_int])
                                        ws_orig_export.cell(row=r, column=max_c + 2, value=ups_map[line_int])
                                except ValueError:
                                    pass
                        
                        wb_orig_export.save(data_filename)
                        
                        # 4. MOVE ORIGINAL SOURCE FILE TO ARCHIVE WITH FORMAT: {filename}_{mo}_{seq}.xlsx
                        if input_method == "Read from Local Directory" and os.path.exists(file_path_or_upload):
                            try:
                                src_dir = os.path.dirname(file_path_or_upload)
                                archive_dir = os.path.join(src_dir, "archive")
                                os.makedirs(archive_dir, exist_ok=True)
                                
                                orig_filename = os.path.basename(file_path_or_upload)
                                name_wo_ext, ext = os.path.splitext(orig_filename)
                                
                                seq = 1
                                while True:
                                    archive_filename = f"{name_wo_ext}_{mo_number}_{seq}{ext}"
                                    dest_path = os.path.join(archive_dir, archive_filename)
                                    if not os.path.exists(dest_path):
                                        break
                                    seq += 1
                                    
                                shutil.move(file_path_or_upload, dest_path)
                            except Exception as arc_err:
                                st.warning(f"File exported successfully but failed to move to archive: {arc_err}")

                        st.success(f"All files exported & archived successfully!\n- MO Report Folder: `{mo_folder_path}`\n- Excel Data Export: `{output_dir_data}`")
                    except Exception as e:
                        st.error(f"An error occurred while exporting files: {e}")

                st.download_button(
                    label="📥 Download Complete Batching Result (CSV)",
                    data=final_batch_editor.to_csv(index=False).encode('utf-8'),
                    file_name="optimized_batching_result.csv",
                    mime="text/css"
                )
        else:
            st.error("⚠️ Failed to match column mapping. Please verify your column letters (A-Z) and header row index.")

    except Exception as x:
        st.error(f"An error occurred while processing the file: {x}")
else:
    st.warning("⚠️ Please select or upload an Excel file via the sidebar on the left.")